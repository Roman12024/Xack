import io, os, json, time, uuid, datetime, secrets, smtplib
from typing import Optional, List
from fastapi import FastAPI, Depends, HTTPException, Header, UploadFile, File, Form, Query, Request
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from .db import get_conn, init_db, now, UPLOAD_DIR, OUTBOX_DIR, DEFAULT_WORKFLOW
from .security import hash_password, verify_password, make_token, verify_token

app = FastAPI(title="ИТ Школа Ростелеком — CRM контроля взаимодействия с вузами",
              description="Система контроля и обработки статистических данных по обучению студентов вузов и школ по ИТ-направлениям",
              version="1.0.0")

init_db()
os.makedirs(UPLOAD_DIR, exist_ok=True)

ALLOWED_EXT = {"png", "jpeg", "jpg", "pdf", "zip", "gzip", "gz", "rar", "doc", "docx", "xls", "xlsx"}

# ---------------- auth ----------------

class LoginIn(BaseModel):
    username: str
    password: str

def current_user(authorization: str = Header(None)):
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(401, "Требуется авторизация")
    payload = verify_token(authorization[7:])
    if not payload:
        raise HTTPException(401, "Недействительный или просроченный токен")
    return payload

def require_role(*roles):
    def dep(user=Depends(current_user)):
        if user["role"] not in roles:
            raise HTTPException(403, "Недостаточно прав (нужна роль: " + ", ".join(roles) + ")")
        return user
    return dep

def parse_user_agent(ua):
    """Упрощённый разбор User-Agent: устройство (модель телефона), ОС, браузер."""
    ua = ua or ""
    low = ua.lower()
    device, os_name, browser = "PC / ноутбук", "Другое", "Другой браузер"
    mobile = 0
    if "iphone" in low:
        device, os_name, mobile = "iPhone", "iOS", 1
    elif "android" in low:
        mobile = 1
        import re as _re
        m = _re.search(r"android [\d.]+;? *([^;)]+)", low)
        device = (m.group(1).strip().title() if m else "Android-устройство")
        os_name = "Android"
    elif "ipad" in low:
        device, os_name, mobile = "iPad", "iPadOS", 1
    elif "windows" in low:
        os_name = "Windows"
    elif "mac os" in low or "macintosh" in low:
        os_name = "macOS"
    elif "linux" in low:
        os_name = "Linux"
    if "edg/" in low:
        browser = "Edge"
    elif "firefox" in low:
        browser = "Firefox"
    elif "chrome" in low:
        browser = "Chrome"
    elif "safari" in low:
        browser = "Safari"
    elif "yabrowser" in low:
        browser = "Яндекс.Браузер"
    return device, os_name, browser, mobile

@app.post("/api/auth/login")
def login(body: LoginIn, request: Request):
    conn = get_conn()
    row = conn.execute("SELECT * FROM users WHERE username=?", (body.username,)).fetchone()
    conn.close()
    if not row or not verify_password(body.password, row["password_hash"]):
        raise HTTPException(401, "Неверный логин или пароль")
    if "is_verified" in row.keys() and not row["is_verified"]:
        raise HTTPException(403, "Подтвердите адрес электронной почты — перейдите по ссылке из письма")
    # --- сессия с телеметрией ---
    ip = request.client.host if request.client else ""
    ua = request.headers.get("user-agent", "")
    device, os_name, browser, mobile = parse_user_agent(ua)
    ts = now()
    conn = get_conn()
    cur = conn.execute(
        "INSERT INTO sessions(user_id,ip,user_agent,device,os,browser,is_mobile,last_seen_at,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
        (row["id"], ip, ua[:400], device, os_name, browser, mobile, ts, ts))
    conn.execute("UPDATE users SET last_login_at=? WHERE id=?", (ts, row["id"]))
    conn.commit()
    sid = cur.lastrowid
    conn.close()
    token = make_token(row["id"], row["username"], row["role"], sid=sid)
    return {"token": token, "role": row["role"], "full_name": row["full_name"], "user_id": row["id"]}


# ---------------- registration & email confirmation ----------------

PUBLIC_URL = os.environ.get("RTK_PUBLIC_URL", "http://localhost:8000")
SMTP_HOST = os.environ.get("RTK_SMTP_HOST", "")
SMTP_PORT = int(os.environ.get("RTK_SMTP_PORT", "465"))
SMTP_USER = os.environ.get("RTK_SMTP_USER", "")
SMTP_PASS = os.environ.get("RTK_SMTP_PASS", "")
SMTP_FROM = os.environ.get("RTK_SMTP_FROM", SMTP_USER or "crm@rt.ru")

def send_confirmation_email(to_email, full_name, token):
    """Отправляет письмо подтверждения. Если SMTP не настроен — письмо сохраняется в outbox (dev-режим)."""
    link = PUBLIC_URL + "/#confirm=" + token
    subject = "ИТ Школа Ростелеком — подтверждение регистрации"
    body = ("Здравствуйте, " + full_name + "!\n\n"
            "Вы зарегистрировались в CRM ИТ Школы Ростелеком.\n"
            "Для подтверждения адреса электронной почты перейдите по ссылке:\n"
            + link + "\n\n"
            "Ссылка действительна до завершения регистрации.\n\n"
            "Если вы не регистрировались — просто проигнорируйте это письмо.")
    if SMTP_HOST:
        msg = ("From: " + SMTP_FROM + "\nTo: " + to_email + "\nSubject: " + subject + "\nContent-Type: text/plain; charset=utf-8\n\n" + body)
        with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT) as s:
            if SMTP_USER:
                s.login(SMTP_USER, SMTP_PASS)
            s.sendmail(SMTP_FROM, [to_email], msg.encode("utf-8"))
        return "sent"
    os.makedirs(OUTBOX_DIR, exist_ok=True)
    fname = os.path.join(OUTBOX_DIR, token + ".eml")
    with open(fname, "w", encoding="utf-8") as f:
        f.write("From: " + SMTP_FROM + "\nTo: " + to_email + "\nSubject: " + subject + "\n\n" + body)
    return "outbox"

class RegisterIn(BaseModel):
    full_name: str
    email: str
    password: str
    agree: bool = False

@app.post("/api/auth/register", status_code=201)
def register(body: RegisterIn):
    email = body.email.strip().lower()
    if "@" not in email or "." not in email.split("@")[-1]:
        raise HTTPException(422, "Некорректный адрес электронной почты")
    if len(body.password) < 6:
        raise HTTPException(422, "Пароль должен быть не короче 6 символов")
    if not body.agree:
        raise HTTPException(422, "Необходимо принять пользовательское соглашение")
    token = secrets.token_urlsafe(32)
    conn = get_conn()
    exist = conn.execute("SELECT id FROM users WHERE username=? OR email=?", (email, email)).fetchone()
    if exist:
        conn.close()
        raise HTTPException(409, "Пользователь с таким e-mail уже зарегистрирован")
    cur = conn.execute(
        "INSERT INTO users(username,password_hash,full_name,role,email,is_verified,confirm_token) VALUES(?,?,?,?,?,0,?)",
        (email, hash_password(body.password), body.full_name.strip(), "user", email, token))
    conn.commit()
    conn.close()
    where = send_confirmation_email(email, body.full_name.strip(), token)
    result = {"ok": True, "message": "Регистрация выполнена. Проверьте почту (" + email + ") и перейдите по ссылке для подтверждения."}
    if where == "outbox":
        # dev-режим: SMTP не настроен, ссылка продублирована в ответе
        result["dev_confirm_link"] = PUBLIC_URL + "/#confirm=" + token
    return result

@app.get("/api/auth/confirm")
def confirm(token: str):
    conn = get_conn()
    row = conn.execute("SELECT id, email FROM users WHERE confirm_token=?", (token,)).fetchone()
    if not row:
        conn.close()
        raise HTTPException(404, "Ссылка подтверждения недействительна или уже использована")
    conn.execute("UPDATE users SET is_verified=1, confirm_token=NULL WHERE id=?", (row["id"],))
    conn.commit()
    conn.close()
    return {"ok": True, "message": "E-mail подтверждён. Теперь можно войти в систему."}

@app.get("/api/auth/me")
def me(user=Depends(current_user)):
    conn = get_conn()
    row = conn.execute("SELECT id,username,full_name,role FROM users WHERE id=?", (user["sub"],)).fetchone()
    conn.close()
    if not row:
        raise HTTPException(401)
    return dict(row)

# ---------------- users ----------------

class UserIn(BaseModel):
    username: str
    password: str
    full_name: str
    role: str = "user"

class UserUpdate(BaseModel):
    full_name: Optional[str] = None
    role: Optional[str] = None
    password: Optional[str] = None

@app.get("/api/users")
def list_users(user=Depends(require_role("manager", "admin"))):
    conn = get_conn()
    rows = conn.execute("SELECT id,username,full_name,role FROM users ORDER BY id").fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.post("/api/users", status_code=201)
def create_user(body: UserIn, user=Depends(require_role("admin"))):
    if body.role not in ("user", "manager", "admin"):
        raise HTTPException(422, "Роль должна быть user, manager или admin")
    conn = get_conn()
    try:
        cur = conn.execute("INSERT INTO users(username,password_hash,full_name,role) VALUES(?,?,?,?)",
                           (body.username, hash_password(body.password), body.full_name, body.role))
        conn.commit()
        uid = cur.lastrowid
    except Exception:
        conn.close()
        raise HTTPException(409, "Пользователь с таким логином уже существует")
    conn.close()
    return {"id": uid, "username": body.username, "full_name": body.full_name, "role": body.role}

@app.put("/api/users/{uid}")
def update_user(uid: int, body: UserUpdate, user=Depends(require_role("admin"))):
    conn = get_conn()
    row = conn.execute("SELECT id FROM users WHERE id=?", (uid,)).fetchone()
    if not row:
        conn.close()
        raise HTTPException(404, "Пользователь не найден")
    if body.full_name:
        conn.execute("UPDATE users SET full_name=? WHERE id=?", (body.full_name, uid))
    if body.role:
        if body.role not in ("user", "manager", "admin"):
            conn.close()
            raise HTTPException(422, "Некорректная роль")
        conn.execute("UPDATE users SET role=? WHERE id=?", (body.role, uid))
    if body.password:
        conn.execute("UPDATE users SET password_hash=? WHERE id=?", (hash_password(body.password), uid))
    conn.commit()
    conn.close()
    return {"ok": True}

# ---------------- catalogs ----------------

class UniversityIn(BaseModel):
    name: str
    vendor: Optional[str] = ""
    software: Optional[str] = ""
    contract_number: Optional[str] = ""
    license_signed: Optional[str] = ""
    license_years: Optional[str] = ""
    transfer_status: Optional[str] = ""
    manager_name: Optional[str] = ""
    uni_responsible: Optional[str] = ""
    comment: Optional[str] = ""

class NameIn(BaseModel):
    name: str
    direction_id: Optional[int] = None
    description: Optional[str] = ""

class ProductIn(BaseModel):
    name: str
    direction_id: Optional[int] = None
    description: Optional[str] = ""

@app.get("/api/universities")
def list_universities(q: Optional[str] = None):
    conn = get_conn()
    if q:
        rows = conn.execute("SELECT * FROM universities WHERE name LIKE ? ORDER BY name", ("%" + q + "%",)).fetchall()
    else:
        rows = conn.execute("SELECT * FROM universities ORDER BY name").fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.post("/api/universities", status_code=201)
def create_university(body: UniversityIn, user=Depends(require_role("manager", "admin"))):
    conn = get_conn()
    try:
        cur = conn.execute("INSERT INTO universities(name,vendor,software,contract_number,license_signed,license_years,transfer_status,manager_name,uni_responsible,comment) VALUES(?,?,?,?,?,?,?,?,?,?)",
                           (body.name, body.vendor, body.software, body.contract_number, body.license_signed, body.license_years, body.transfer_status, body.manager_name, body.uni_responsible, body.comment))
        conn.commit()
        uid = cur.lastrowid
    except Exception:
        conn.close()
        raise HTTPException(409, "ВУЗ с таким названием уже есть в каталоге")
    conn.close()
    return {"id": uid}

@app.put("/api/universities/{uid}")
def update_university(uid: int, body: UniversityIn, user=Depends(require_role("manager", "admin"))):
    conn = get_conn()
    if not conn.execute("SELECT id FROM universities WHERE id=?", (uid,)).fetchone():
        conn.close()
        raise HTTPException(404, "ВУЗ не найден")
    conn.execute("UPDATE universities SET name=?,vendor=?,software=?,contract_number=?,license_signed=?,license_years=?,transfer_status=?,manager_name=?,uni_responsible=?,comment=? WHERE id=?",
                 (body.name, body.vendor, body.software, body.contract_number, body.license_signed, body.license_years, body.transfer_status, body.manager_name, body.uni_responsible, body.comment, uid))
    conn.commit()
    conn.close()
    return {"ok": True}

@app.delete("/api/universities/{uid}")
def delete_university(uid: int, user=Depends(require_role("manager", "admin"))):
    conn = get_conn()
    used = conn.execute("SELECT COUNT(*) c FROM interactions WHERE university_id=?", (uid,)).fetchone()["c"]
    if used:
        conn.close()
        raise HTTPException(409, "Нельзя удалить: существуют взаимодействия с этим вузом")
    conn.execute("DELETE FROM universities WHERE id=?", (uid,))
    conn.commit()
    conn.close()
    return {"ok": True}

@app.get("/api/directions")
def list_directions():
    conn = get_conn()
    rows = conn.execute("SELECT * FROM directions ORDER BY name").fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.post("/api/directions", status_code=201)
def create_direction(body: NameIn, user=Depends(require_role("manager", "admin"))):
    conn = get_conn()
    try:
        cur = conn.execute("INSERT INTO directions(name) VALUES(?)", (body.name,))
        conn.commit()
        did = cur.lastrowid
    except Exception:
        conn.close()
        raise HTTPException(409, "Направление уже существует")
    conn.close()
    return {"id": did}

@app.delete("/api/directions/{did}")
def delete_direction(did: int, user=Depends(require_role("manager", "admin"))):
    conn = get_conn()
    used = conn.execute("SELECT COUNT(*) c FROM products WHERE direction_id=?", (did,)).fetchone()["c"]
    if used:
        conn.close()
        raise HTTPException(409, "Нельзя удалить: есть ИТ-продукты этого направления")
    conn.execute("DELETE FROM directions WHERE id=?", (did,))
    conn.commit()
    conn.close()
    return {"ok": True}

@app.get("/api/products")
def list_products(direction_id: Optional[int] = None):
    conn = get_conn()
    if direction_id:
        rows = conn.execute("SELECT p.*, d.name AS direction FROM products p LEFT JOIN directions d ON d.id=p.direction_id WHERE p.direction_id=? ORDER BY p.name", (direction_id,)).fetchall()
    else:
        rows = conn.execute("SELECT p.*, d.name AS direction FROM products p LEFT JOIN directions d ON d.id=p.direction_id ORDER BY p.name").fetchall()
    conn.close()
    out = []
    for r in rows:
        d = dict(r)
        d["photo_url"] = "/api/products/" + str(d["id"]) + "/photo" if d.get("photo") else None
        out.append(d)
    return out

@app.get("/api/products/{pid}")
def get_product(pid: int, user=Depends(current_user)):
    conn = get_conn()
    row = conn.execute("SELECT p.*, d.name AS direction FROM products p LEFT JOIN directions d ON d.id=p.direction_id WHERE p.id=?", (pid,)).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "ИТ-проект не найден")
    d = dict(row)
    d["photo_url"] = "/api/products/" + str(d["id"]) + "/photo" if d.get("photo") else None
    return d

@app.post("/api/products", status_code=201)
def create_product(name: str = Form(...), direction_id: Optional[str] = Form(None),
                   description: str = Form(""), requirements: str = Form(""),
                   start_date: str = Form(""), end_date: str = Form(""),
                   photo: Optional[UploadFile] = File(None),
                   user=Depends(require_role("manager", "admin"))):
    photo_name = save_course_photo(photo) if photo and photo.filename else ""
    did = int(direction_id) if direction_id else None
    conn = get_conn()
    cur = conn.execute("INSERT INTO products(name,direction_id,description,requirements,start_date,end_date,photo) VALUES(?,?,?,?,?,?,?)",
                       (name.strip(), did, description, requirements, start_date, end_date, photo_name))
    conn.commit()
    pid = cur.lastrowid
    conn.close()
    return {"id": pid}

@app.put("/api/products/{pid}")
def update_product(pid: int, name: str = Form(...), direction_id: Optional[str] = Form(None),
                   description: str = Form(""), requirements: str = Form(""),
                   start_date: str = Form(""), end_date: str = Form(""),
                   photo: Optional[UploadFile] = File(None),
                   user=Depends(require_role("manager", "admin"))):
    conn = get_conn()
    row = conn.execute("SELECT photo FROM products WHERE id=?", (pid,)).fetchone()
    if not row:
        conn.close()
        raise HTTPException(404, "ИТ-проект не найден")
    photo_name = row["photo"]
    if photo and photo.filename:
        photo_name = save_course_photo(photo)
    did = int(direction_id) if direction_id else None
    conn.execute("UPDATE products SET name=?,direction_id=?,description=?,requirements=?,start_date=?,end_date=?,photo=? WHERE id=?",
                 (name.strip(), did, description, requirements, start_date, end_date, photo_name, pid))
    conn.commit()
    conn.close()
    return {"ok": True}

@app.get("/api/products/{pid}/photo")
def product_photo(pid: int, token: Optional[str] = Query(None), authorization: Optional[str] = Header(None)):
    current_user(authorization or ("Bearer " + token if token else None))
    conn = get_conn()
    row = conn.execute("SELECT photo FROM products WHERE id=?", (pid,)).fetchone()
    conn.close()
    if not row or not row["photo"]:
        raise HTTPException(404, "Фото не найдено")
    path = os.path.join(UPLOAD_DIR, row["photo"])
    if not os.path.exists(path):
        raise HTTPException(404, "Файл фото отсутствует на диске")
    return FileResponse(path)

@app.delete("/api/products/{pid}")
def delete_product(pid: int, user=Depends(require_role("manager", "admin"))):
    conn = get_conn()
    used = conn.execute("SELECT COUNT(*) c FROM interactions WHERE product_id=?", (pid,)).fetchone()["c"]
    if used:
        conn.close()
        raise HTTPException(409, "Нельзя удалить: есть взаимодействия с этим продуктом")
    conn.execute("DELETE FROM products WHERE id=?", (pid,))
    conn.commit()
    conn.close()
    return {"ok": True}

# ---------------- import catalog from xlsx ----------------

IMPORT_FIELDS = {
    "название вуза": "name", "вендор": "vendor", "по": "software",
    "номер договора": "contract_number", "подписание лицензии": "license_signed",
    "срок действия лицензии": "license_years", "статус по передачи": "transfer_status",
    "фио менеджера": "manager_name", "ответственные от вуза": "uni_responsible",
    "комментарий": "comment",
}

@app.post("/api/catalogs/import")
def import_catalog(file: UploadFile = File(...), user=Depends(require_role("manager", "admin"))):
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file.file.read()), read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception:
        raise HTTPException(422, "Не удалось прочитать xlsx-файл")
    if not rows:
        raise HTTPException(422, "Файл пуст")
    header = [str(h).strip().lower() if h else "" for h in rows[0]]
    colmap = {}
    for idx, h in enumerate(header):
        for key, field in IMPORT_FIELDS.items():
            if key in h:
                colmap[idx] = field
    if "name" not in colmap.values():
        raise HTTPException(422, "В файле не найдена колонка 'Название ВУЗа'")
    conn = get_conn()
    created, updated, skipped = 0, 0, 0
    for row in rows[1:]:
        data = {f: "" for f in IMPORT_FIELDS.values()}
        for idx, field in colmap.items():
            if idx < len(row) and row[idx] is not None:
                data[field] = str(row[idx]).strip()
        if not data["name"]:
            skipped += 1
            continue
        exist = conn.execute("SELECT id FROM universities WHERE name=?", (data["name"],)).fetchone()
        if exist:
            conn.execute("UPDATE universities SET vendor=?,software=?,contract_number=?,license_signed=?,license_years=?,transfer_status=?,manager_name=?,uni_responsible=?,comment=? WHERE id=?",
                         (data["vendor"], data["software"], data["contract_number"], data["license_signed"], data["license_years"], data["transfer_status"], data["manager_name"], data["uni_responsible"], data["comment"], exist["id"]))
            updated += 1
        else:
            conn.execute("INSERT INTO universities(name,vendor,software,contract_number,license_signed,license_years,transfer_status,manager_name,uni_responsible,comment) VALUES(?,?,?,?,?,?,?,?,?,?)",
                         (data["name"], data["vendor"], data["software"], data["contract_number"], data["license_signed"], data["license_years"], data["transfer_status"], data["manager_name"], data["uni_responsible"], data["comment"]))
            created += 1
    conn.commit()
    conn.close()
    return {"created": created, "updated": updated, "skipped": skipped}

# ---------------- workflows ----------------

class WorkflowIn(BaseModel):
    name: str
    statuses: List[str] = []

class StatusIn(BaseModel):
    name: str

class StatusUpdate(BaseModel):
    name: Optional[str] = None
    position: Optional[int] = None

@app.get("/api/workflows")
def list_workflows(user=Depends(current_user)):
    conn = get_conn()
    rows = conn.execute("SELECT * FROM workflows ORDER BY id").fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.post("/api/workflows", status_code=201)
def create_workflow(body: WorkflowIn, user=Depends(require_role("admin"))):
    if not body.statuses:
        body.statuses = list(DEFAULT_WORKFLOW)
    conn = get_conn()
    cur = conn.execute("INSERT INTO workflows(name,is_default) VALUES(?,0)", (body.name,))
    wf = cur.lastrowid
    for pos, name in enumerate(body.statuses):
        conn.execute("INSERT INTO workflow_statuses(workflow_id,name,position) VALUES(?,?,?)", (wf, name, pos))
    conn.commit()
    conn.close()
    return {"id": wf}

@app.get("/api/workflows/{wid}/statuses")
def list_statuses(wid: int, user=Depends(current_user)):
    conn = get_conn()
    rows = conn.execute("SELECT * FROM workflow_statuses WHERE workflow_id=? ORDER BY position", (wid,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.post("/api/workflows/{wid}/statuses", status_code=201)
def add_status(wid: int, body: StatusIn, user=Depends(require_role("admin"))):
    conn = get_conn()
    maxpos = conn.execute("SELECT COALESCE(MAX(position),-1) m FROM workflow_statuses WHERE workflow_id=?", (wid,)).fetchone()["m"]
    cur = conn.execute("INSERT INTO workflow_statuses(workflow_id,name,position) VALUES(?,?,?)", (wid, body.name, maxpos + 1))
    conn.commit()
    sid = cur.lastrowid
    conn.close()
    return {"id": sid}

@app.put("/api/statuses/{sid}")
def update_status(sid: int, body: StatusUpdate, user=Depends(require_role("admin"))):
    conn = get_conn()
    if not conn.execute("SELECT id FROM workflow_statuses WHERE id=?", (sid,)).fetchone():
        conn.close()
        raise HTTPException(404, "Статус не найден")
    if body.name:
        conn.execute("UPDATE workflow_statuses SET name=? WHERE id=?", (body.name, sid))
    if body.position is not None:
        conn.execute("UPDATE workflow_statuses SET position=? WHERE id=?", (body.position, sid))
    conn.commit()
    conn.close()
    return {"ok": True}

# ---------------- interactions ----------------

class InteractionIn(BaseModel):
    university_id: int
    product_id: Optional[int] = None
    workflow_id: Optional[int] = None
    responsible_id: Optional[int] = None

class TransitionIn(BaseModel):
    to_status_id: int
    comment: str = ""

class ResponsibleIn(BaseModel):
    user_id: Optional[int] = None

def build_filter(from_date, to_date, university_id, direction_id, product_id, status_id, responsible_id, params):
    where, p = [], []
    if from_date:
        where.append("date(i.updated_at) >= date(?)"); p.append(from_date)
    if to_date:
        where.append("date(i.updated_at) <= date(?)"); p.append(to_date)
    if university_id:
        where.append("i.university_id = ?"); p.append(university_id)
    if direction_id:
        where.append("p.direction_id = ?"); p.append(direction_id)
    if product_id:
        where.append("i.product_id = ?"); p.append(product_id)
    if status_id:
        where.append("i.current_status_id = ?"); p.append(status_id)
    if responsible_id:
        where.append("i.responsible_id = ?"); p.append(responsible_id)
    return (" WHERE " + " AND ".join(where)) if where else "", p

BASE_SQL = """
    SELECT i.id, i.university_id, i.product_id, i.workflow_id, i.current_status_id, i.responsible_id,
           i.created_at, i.updated_at,
           u.name AS university, p.name AS product, d.name AS direction,
           ws.name AS status, ws.position AS status_position,
           ru.full_name AS responsible
    FROM interactions i
    JOIN universities u ON u.id = i.university_id
    LEFT JOIN products p ON p.id = i.product_id
    LEFT JOIN directions d ON d.id = p.direction_id
    LEFT JOIN workflow_statuses ws ON ws.id = i.current_status_id
    LEFT JOIN users ru ON ru.id = i.responsible_id
"""

@app.get("/api/interactions")
def list_interactions(from_date: Optional[str] = Query(None), to_date: Optional[str] = Query(None),
                      university_id: Optional[int] = None, direction_id: Optional[int] = None,
                      product_id: Optional[int] = None, status_id: Optional[int] = None,
                      responsible_id: Optional[int] = None,
                      user=Depends(current_user)):
    where, p = build_filter(from_date, to_date, university_id, direction_id, product_id, status_id, responsible_id, None)
    conn = get_conn()
    rows = conn.execute(BASE_SQL + where + " ORDER BY i.updated_at DESC", p).fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.post("/api/interactions", status_code=201)
def create_interaction(body: InteractionIn, user=Depends(current_user)):
    conn = get_conn()
    if not conn.execute("SELECT id FROM universities WHERE id=?", (body.university_id,)).fetchone():
        conn.close()
        raise HTTPException(404, "ВУЗ не найден")
    wf = body.workflow_id or (conn.execute("SELECT id FROM workflows WHERE is_default=1").fetchone() or [None])[0]
    if isinstance(wf, int):
        wf_id = wf
    else:
        wf_id = wf["id"] if wf else None
    if not wf_id:
        conn.close()
        raise HTTPException(422, "Нет доступного workflow")
    first = conn.execute("SELECT id FROM workflow_statuses WHERE workflow_id=? ORDER BY position LIMIT 1", (wf_id,)).fetchone()
    ts = now()
    cur = conn.execute("INSERT INTO interactions(university_id,product_id,workflow_id,current_status_id,responsible_id,created_at,updated_at) VALUES(?,?,?,?,?,?,?)",
                       (body.university_id, body.product_id, wf_id, first["id"] if first else None, body.responsible_id or user["sub"], ts, ts))
    iid = cur.lastrowid
    conn.execute("INSERT INTO interaction_history(interaction_id,to_status_id,to_status_name,comment,user_id,created_at) VALUES(?,?,?,?,?,?)",
                 (iid, first["id"] if first else None, "Создано взаимодействие", "", user["sub"], ts))
    conn.commit()
    conn.close()
    return {"id": iid}

@app.get("/api/interactions/{iid}")
def get_interaction(iid: int, user=Depends(current_user)):
    conn = get_conn()
    row = conn.execute(BASE_SQL + " WHERE i.id=?", (iid,)).fetchone()
    if not row:
        conn.close()
        raise HTTPException(404, "Взаимодействие не найдено")
    hist = conn.execute("SELECT h.*, u.full_name AS user_name FROM interaction_history h LEFT JOIN users u ON u.id=h.user_id WHERE h.interaction_id=? ORDER BY h.created_at, h.id", (iid,)).fetchall()
    atts = conn.execute("SELECT a.id, a.history_id, a.filename, a.size, a.created_at FROM attachments a JOIN interaction_history h ON h.id=a.history_id WHERE h.interaction_id=? ORDER BY a.id", (iid,)).fetchall()
    statuses = conn.execute("SELECT id,name,position FROM workflow_statuses WHERE workflow_id=? ORDER BY position", (row["workflow_id"],)).fetchall()
    conn.close()
    d = dict(row)
    d["history"] = [dict(h) for h in hist]
    d["attachments"] = [dict(a) for a in atts]
    d["available_statuses"] = [dict(s) for s in statuses]
    return d

@app.post("/api/interactions/{iid}/transition")
def transition(iid: int, body: TransitionIn, user=Depends(current_user)):
    conn = get_conn()
    inter = conn.execute("SELECT * FROM interactions WHERE id=?", (iid,)).fetchone()
    if not inter:
        conn.close()
        raise HTTPException(404, "Взаимодействие не найдено")
    st = conn.execute("SELECT * FROM workflow_statuses WHERE id=? AND workflow_id=?", (body.to_status_id, inter["workflow_id"])).fetchone()
    if not st:
        conn.close()
        raise HTTPException(422, "Статус не принадлежит workflow этого взаимодействия")
    cur_st = conn.execute("SELECT name FROM workflow_statuses WHERE id=?", (inter["current_status_id"],)).fetchone()
    ts = now()
    cur = conn.execute("INSERT INTO interaction_history(interaction_id,from_status_id,from_status_name,to_status_id,to_status_name,comment,user_id,created_at) VALUES(?,?,?,?,?,?,?,?)",
                       (iid, inter["current_status_id"], cur_st["name"] if cur_st else None, st["id"], st["name"], body.comment, user["sub"], ts))
    conn.execute("UPDATE interactions SET current_status_id=?, updated_at=? WHERE id=?", (st["id"], ts, iid))
    conn.commit()
    conn.close()
    return {"history_id": cur.lastrowid, "status": st["name"]}

@app.put("/api/interactions/{iid}/responsible")
def set_responsible(iid: int, body: ResponsibleIn, user=Depends(require_role("manager", "admin"))):
    conn = get_conn()
    if not conn.execute("SELECT id FROM interactions WHERE id=?", (iid,)).fetchone():
        conn.close()
        raise HTTPException(404, "Взаимодействие не найдено")
    if body.user_id and not conn.execute("SELECT id FROM users WHERE id=?", (body.user_id,)).fetchone():
        conn.close()
        raise HTTPException(404, "Пользователь не найден")
    conn.execute("UPDATE interactions SET responsible_id=?, updated_at=? WHERE id=?", (body.user_id, now(), iid))
    conn.commit()
    conn.close()
    return {"ok": True}

@app.post("/api/interactions/{iid}/attachments")
def upload_attachment(iid: int, file: UploadFile = File(...), user=Depends(current_user)):
    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in ALLOWED_EXT:
        raise HTTPException(422, "Недопустимый формат файла. Разрешены: " + ", ".join(sorted(ALLOWED_EXT)))
    conn = get_conn()
    inter = conn.execute("SELECT id, current_status_id FROM interactions WHERE id=?", (iid,)).fetchone()
    if not inter:
        conn.close()
        raise HTTPException(404, "Взаимодействие не найдено")
    st = conn.execute("SELECT name FROM workflow_statuses WHERE id=?", (inter["current_status_id"],)).fetchone()
    ts = now()
    stored = uuid.uuid4().hex + "." + ext
    content = file.file.read()
    if len(content) > 20 * 1024 * 1024:
        conn.close()
        raise HTTPException(413, "Файл больше 20 МБ")
    with open(os.path.join(UPLOAD_DIR, stored), "wb") as f:
        f.write(content)
    cur = conn.execute("INSERT INTO interaction_history(interaction_id,from_status_id,from_status_name,to_status_id,to_status_name,comment,user_id,created_at) VALUES(?,?,?,?,?,?,?,?)",
                       (iid, inter["current_status_id"], st["name"] if st else None, inter["current_status_id"], st["name"] if st else None, "Прикреплен файл: " + file.filename, user["sub"], ts))
    conn.execute("INSERT INTO attachments(history_id,filename,stored_name,size,created_at) VALUES(?,?,?,?,?)",
                 (cur.lastrowid, file.filename, stored, len(content), ts))
    conn.commit()
    conn.close()
    return {"ok": True, "filename": file.filename}

@app.get("/api/attachments/{aid}/download")
def download_attachment(aid: int, token: Optional[str] = Query(None), authorization: Optional[str] = Header(None)):
    auth = authorization or ("Bearer " + token if token else None)
    current_user(auth)
    conn = get_conn()
    row = conn.execute("SELECT * FROM attachments WHERE id=?", (aid,)).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404, "Файл не найден")
    path = os.path.join(UPLOAD_DIR, row["stored_name"])
    if not os.path.exists(path):
        raise HTTPException(404, "Файл отсутствует на диске")
    return FileResponse(path, filename=row["filename"])

# ---------------- reports ----------------

REPORT_COLUMNS = ["university", "direction", "product", "status", "responsible"]

def query_report(from_date, to_date, university_id, direction_id, product_id, status_id, responsible_id):
    where, p = build_filter(from_date, to_date, university_id, direction_id, product_id, status_id, responsible_id, None)
    conn = get_conn()
    rows = conn.execute(BASE_SQL + where + " ORDER BY u.name, p.name", p).fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.get("/api/reports/stats")
def report_stats(from_date: Optional[str] = None, to_date: Optional[str] = None,
                 university_id: Optional[int] = None, direction_id: Optional[int] = None,
                 product_id: Optional[int] = None, status_id: Optional[int] = None,
                 responsible_id: Optional[int] = None,
                 token: Optional[str] = Query(None), authorization: Optional[str] = Header(None)):
    _payload = current_user(authorization or ("Bearer " + token if token else None))
    if _payload["role"] == "user":
        raise HTTPException(403, "Статистика доступна только руководителям и администраторам")
    rows = query_report(from_date, to_date, university_id, direction_id, product_id, status_id, responsible_id)
    by_status, by_university, by_responsible = {}, {}, {}
    for r in rows:
        by_status[r["status"] or "Без статуса"] = by_status.get(r["status"] or "Без статуса", 0) + 1
        by_university[r["university"]] = by_university.get(r["university"], 0) + 1
        by_responsible[r["responsible"] or "Не назначен"] = by_responsible.get(r["responsible"] or "Не назначен", 0) + 1
    return {"total": len(rows), "by_status": by_status, "by_university": by_university, "by_responsible": by_responsible}

@app.get("/api/reports/export")
def report_export(fmt: str = "xlsx", from_date: Optional[str] = None, to_date: Optional[str] = None,
                  university_id: Optional[int] = None, direction_id: Optional[int] = None,
                  product_id: Optional[int] = None, status_id: Optional[int] = None,
                  responsible_id: Optional[int] = None,
                  token: Optional[str] = Query(None), authorization: Optional[str] = Header(None)):
    _payload = current_user(authorization or ("Bearer " + token if token else None))
    if _payload["role"] == "user":
        raise HTTPException(403, "Отчёты доступны только руководителям и администраторам")
    rows = query_report(from_date, to_date, university_id, direction_id, product_id, status_id, responsible_id)
    headers = ["Наименование ВУЗа", "ИТ-направление", "ИТ-продукт", "Статус работы с вузом", "Ответственный"]
    data = [[r["university"], r["direction"] or "", r["product"] or "", r["status"] or "", r["responsible"] or ""] for r in rows]
    if fmt == "xlsx":
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Отчет"
        ws.append(["Отчет по взаимодействиям с вузами (ИТ Школа Ростелеком)"])
        ws.append(["Период:", (from_date or "…") + " — " + (to_date or "…"), "Сформирован:", now()])
        ws.append([])
        ws.append(headers)
        for c in ws[4]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="7700FF")
        for row in data:
            ws.append(row)
        for i, w in enumerate([34, 20, 28, 44, 22], start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": "attachment; filename=report.xlsx"})
    elif fmt == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4), title="Отчет по взаимодействиям")
        styles = getSampleStyleSheet()
        story = [Paragraph("Отчет по взаимодействиям с вузами (ИТ Школа Ростелеком)", styles["Title"]),
                 Paragraph("Период: " + (from_date or "…") + " — " + (to_date or "…") + ", сформирован: " + now(), styles["Normal"]),
                 Spacer(1, 12)]
        table_data = [headers] + data
        t = Table(table_data, repeatRows=1, colWidths=[190, 100, 140, 230, 110])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#7700FF")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3E8FF")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(t)
        doc.build(story)
        buf.seek(0)
        return StreamingResponse(buf, media_type="application/pdf",
                                 headers={"Content-Disposition": "attachment; filename=report.pdf"})
    else:
        raise HTTPException(422, "fmt должен быть xlsx или pdf")

@app.get("/api/reports/chart.png")
def report_chart(kind: str = Query("by_status"), from_date: Optional[str] = None, to_date: Optional[str] = None,
                 university_id: Optional[int] = None, direction_id: Optional[int] = None,
                 product_id: Optional[int] = None, status_id: Optional[int] = None,
                 responsible_id: Optional[int] = None,
                 token: Optional[str] = Query(None), authorization: Optional[str] = Header(None)):
    user_payload = current_user(authorization or ("Bearer " + token if token else None))
    if user_payload["role"] == "user":
        raise HTTPException(403, "Диаграммы доступны только руководителям и администраторам")
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    rows = query_report(from_date, to_date, university_id, direction_id, product_id, status_id, responsible_id)
    stats = {"by_status": {}, "by_university": {}, "by_responsible": {}}
    for r in rows:
        stats["by_status"][r["status"] or "Без статуса"] = stats["by_status"].get(r["status"] or "Без статуса", 0) + 1
        stats["by_university"][r["university"]] = stats["by_university"].get(r["university"], 0) + 1
        stats["by_responsible"][r["responsible"] or "Не назначен"] = stats["by_responsible"].get(r["responsible"] or "Не назначен", 0) + 1
    key = kind if kind in stats and isinstance(stats[kind], dict) else "by_status"
    d = stats[key]
    if not d:
        d = {"Нет данных": 1}
    labels = [k if len(k) <= 30 else k[:27] + "..." for k in d.keys()]
    vals = list(d.values())
    fig, ax = plt.subplots(figsize=(9, 4.8), dpi=110)
    bars = ax.barh(range(len(vals)), vals, color="#7700FF")
    ax.set_yticks(range(len(labels)))
    ax.set_yticklabels(labels, fontsize=9)
    ax.invert_yaxis()
    ax.set_title("Взаимодействия с вузами: " + {"by_status": "по статусам", "by_university": "по вузам", "by_responsible": "по ответственным"}.get(key, key))
    ax.set_xlabel("Количество")
    for b, v in zip(bars, vals):
        ax.text(v + 0.05, b.get_y() + b.get_height() / 2, str(v), va="center", fontsize=9)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    buf.seek(0)
    return StreamingResponse(buf, media_type="image/png")



# ---------------- profile (личный кабинет участника) ----------------

PROFILE_FIELDS = ["full_name", "phone", "birth_date", "city", "organization",
                  "position", "education", "skills", "telegram", "about", "portfolio"]

class ProfileIn(BaseModel):
    model_config = {"extra": "forbid"}
    full_name: Optional[str] = None
    phone: Optional[str] = None
    birth_date: Optional[str] = None
    city: Optional[str] = None
    organization: Optional[str] = None
    position: Optional[str] = None
    education: Optional[str] = None
    skills: Optional[str] = None
    telegram: Optional[str] = None
    about: Optional[str] = None
    portfolio: Optional[str] = None

@app.get("/api/profile")
def get_profile(user=Depends(current_user)):
    conn = get_conn()
    row = conn.execute("SELECT * FROM users WHERE id=?", (user["sub"],)).fetchone()
    conn.close()
    if not row:
        raise HTTPException(404)
    d = {k: row[k] for k in row.keys() if k not in ("password_hash", "confirm_token")}
    d["avatar_url"] = "/api/profile/avatar" if d.get("avatar") else None
    return d

@app.put("/api/profile")
def update_profile(body: ProfileIn, user=Depends(current_user)):
    data = body.model_dump(exclude_none=True)
    if not data:
        return {"ok": True}
    for k in data:
        if k not in PROFILE_FIELDS:
            raise HTTPException(422, "Неизвестное поле: " + k)
    sets = ", ".join(k + "=?" for k in data)
    conn = get_conn()
    conn.execute("UPDATE users SET " + sets + " WHERE id=?", list(data.values()) + [user["sub"]])
    conn.commit()
    conn.close()
    return {"ok": True}

@app.post("/api/profile/avatar")
def upload_avatar(file: UploadFile = File(...), user=Depends(current_user)):
    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in COURSE_PHOTO_EXT:
        raise HTTPException(422, "Аватар: разрешены только " + ", ".join(sorted(COURSE_PHOTO_EXT)))
    content = file.file.read()
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(413, "Фото больше 5 МБ")
    stored = uuid.uuid4().hex + "." + ext
    with open(os.path.join(UPLOAD_DIR, stored), "wb") as f:
        f.write(content)
    conn = get_conn()
    old = conn.execute("SELECT avatar FROM users WHERE id=?", (user["sub"],)).fetchone()
    conn.execute("UPDATE users SET avatar=? WHERE id=?", (stored, user["sub"]))
    conn.commit()
    conn.close()
    if old and old["avatar"]:
        try:
            os.remove(os.path.join(UPLOAD_DIR, old["avatar"]))
        except OSError:
            pass
    return {"ok": True, "avatar_url": "/api/profile/avatar"}

@app.get("/api/profile/avatar")
def get_avatar(token: Optional[str] = Query(None), authorization: Optional[str] = Header(None), _: str = ""):
    payload = current_user(authorization or ("Bearer " + token if token else None))
    conn = get_conn()
    row = conn.execute("SELECT avatar FROM users WHERE id=?", (payload["sub"],)).fetchone()
    conn.close()
    if not row or not row["avatar"]:
        raise HTTPException(404, "Аватар не загружен")
    path = os.path.join(UPLOAD_DIR, row["avatar"])
    if not os.path.exists(path):
        raise HTTPException(404, "Файл аватара отсутствует")
    return FileResponse(path)

# ---------------- courses (курсы: сроки, описание, требования, фото) ----------------

COURSE_PHOTO_EXT = {"png", "jpg", "jpeg", "webp", "gif"}

COURSE_BASE_SQL = """
    SELECT c.*, d.name AS direction, u.full_name AS author
    FROM courses c
    LEFT JOIN directions d ON d.id = c.direction_id
    LEFT JOIN users u ON u.id = c.created_by
"""

def save_course_photo(file):
    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in COURSE_PHOTO_EXT:
        raise HTTPException(422, "Фото: разрешены только " + ", ".join(sorted(COURSE_PHOTO_EXT)))
    content = file.file.read()
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(413, "Фото больше 5 МБ")
    stored = uuid.uuid4().hex + "." + ext
    with open(os.path.join(UPLOAD_DIR, stored), "wb") as f:
        f.write(content)
    return stored

@app.get("/api/courses")
def list_courses(user=Depends(current_user)):
    conn = get_conn()
    rows = conn.execute(COURSE_BASE_SQL + " ORDER BY c.id DESC").fetchall()
    conn.close()
    today = datetime.date.today().isoformat()
    out = []
    for r in rows:
        d = dict(r)
        # активный курс: дата окончания не наступила (или не задана)
        d["is_active"] = (not d.get("end_date")) or (d["end_date"] >= today)
        # обычным пользователям — только активные курсы; staff видит все
        if user["role"] == "user" and not d["is_active"]:
            continue
        d["photo_url"] = "/api/courses/" + str(d["id"]) + "/photo" if d.get("photo") else None
        d["description"] = (d["description"] or "")[:180]
        d["requirements"] = (d["requirements"] or "")[:120]
        out.append(d)
    return out

@app.get("/api/courses/{cid}")
def get_course(cid: int, user=Depends(current_user)):
    conn = get_conn()
    row = conn.execute(COURSE_BASE_SQL + " WHERE c.id=?", (cid,)).fetchone()
    if not row:
        conn.close()
        raise HTTPException(404, "Курс не найден")
    d = dict(row)
    d["photo_url"] = "/api/courses/" + str(d["id"]) + "/photo" if d.get("photo") else None
    d["enrolled"] = bool(conn.execute("SELECT 1 FROM enrollments WHERE course_id=? AND user_id=?", (cid, user["sub"])).fetchone())
    d["participants_count"] = conn.execute("SELECT COUNT(*) c FROM enrollments WHERE course_id=?", (cid,)).fetchone()["c"]
    conn.close()
    return d

@app.post("/api/courses/{cid}/enroll")
def enroll_course(cid: int, user=Depends(current_user)):
    conn = get_conn()
    if not conn.execute("SELECT id FROM courses WHERE id=?", (cid,)).fetchone():
        conn.close()
        raise HTTPException(404, "Курс не найден")
    conn.execute("INSERT OR IGNORE INTO enrollments(course_id,user_id,created_at) VALUES(?,?,?)",
                 (cid, user["sub"], now()))
    conn.commit()
    conn.close()
    return {"ok": True}

@app.delete("/api/courses/{cid}/enroll")
def unenroll_course(cid: int, user=Depends(current_user)):
    conn = get_conn()
    conn.execute("DELETE FROM enrollments WHERE course_id=? AND user_id=?", (cid, user["sub"]))
    conn.commit()
    conn.close()
    return {"ok": True}

@app.get("/api/courses/{cid}/participants")
def course_participants(cid: int, user=Depends(require_role("manager", "admin"))):
    conn = get_conn()
    if not conn.execute("SELECT id FROM courses WHERE id=?", (cid,)).fetchone():
        conn.close()
        raise HTTPException(404, "Курс не найден")
    rows = conn.execute(
        "SELECT u.id, u.full_name, u.email, u.phone, u.birth_date, u.city, u.organization, "
        "u.position, u.education, u.skills, u.telegram, u.portfolio, u.about, u.avatar, "
        "e.created_at AS enrolled_at "
        "FROM enrollments e JOIN users u ON u.id=e.user_id WHERE e.course_id=? ORDER BY e.created_at", (cid,)).fetchall()
    conn.close()
    out = []
    for r in rows:
        d = dict(r)
        d["avatar_url"] = "/api/users/" + str(d["id"]) + "/avatar" if d.get("avatar") else None
        out.append(d)
    return out

@app.get("/api/users/{uid}/avatar")
def user_avatar(uid: int, token: Optional[str] = Query(None), authorization: Optional[str] = Header(None)):
    payload = current_user(authorization or ("Bearer " + token if token else None))
    if payload["role"] == "user" and payload["sub"] != uid:
        raise HTTPException(403, "Недостаточно прав")
    conn = get_conn()
    row = conn.execute("SELECT avatar FROM users WHERE id=?", (uid,)).fetchone()
    conn.close()
    if not row or not row["avatar"]:
        raise HTTPException(404, "Аватар не загружен")
    path = os.path.join(UPLOAD_DIR, row["avatar"])
    if not os.path.exists(path):
        raise HTTPException(404, "Файл аватара отсутствует")
    return FileResponse(path)

@app.post("/api/courses", status_code=201)
def create_course(title: str = Form(...), description: str = Form(""),
                  requirements: str = Form(""), start_date: str = Form(""),
                  end_date: str = Form(""), direction_id: Optional[str] = Form(None),
                  photo: Optional[UploadFile] = File(None),
                  user=Depends(require_role("manager", "admin"))):
    conn = get_conn()
    photo_name = ""
    if photo and photo.filename:
        photo_name = save_course_photo(photo)
    did = int(direction_id) if direction_id else None
    cur = conn.execute(
        "INSERT INTO courses(title,direction_id,description,requirements,start_date,end_date,photo,created_by,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
        (title.strip(), did, description, requirements, start_date, end_date, photo_name, user["sub"], now()))
    conn.commit()
    cid = cur.lastrowid
    conn.close()
    return {"id": cid}

@app.put("/api/courses/{cid}")
def update_course(cid: int, title: str = Form(...), description: str = Form(""),
                  requirements: str = Form(""), start_date: str = Form(""),
                  end_date: str = Form(""), direction_id: Optional[str] = Form(None),
                  photo: Optional[UploadFile] = File(None),
                  user=Depends(require_role("manager", "admin"))):
    conn = get_conn()
    row = conn.execute("SELECT * FROM courses WHERE id=?", (cid,)).fetchone()
    if not row:
        conn.close()
        raise HTTPException(404, "Курс не найден")
    photo_name = row["photo"]
    if photo and photo.filename:
        photo_name = save_course_photo(photo)
    did = int(direction_id) if direction_id else None
    conn.execute(
        "UPDATE courses SET title=?,direction_id=?,description=?,requirements=?,start_date=?,end_date=?,photo=? WHERE id=?",
        (title.strip(), did, description, requirements, start_date, end_date, photo_name, cid))
    conn.commit()
    conn.close()
    return {"ok": True}

@app.delete("/api/courses/{cid}")
def delete_course(cid: int, user=Depends(require_role("manager", "admin"))):
    conn = get_conn()
    row = conn.execute("SELECT photo FROM courses WHERE id=?", (cid,)).fetchone()
    if not row:
        conn.close()
        raise HTTPException(404, "Курс не найден")
    if row["photo"]:
        try:
            os.remove(os.path.join(UPLOAD_DIR, row["photo"]))
        except OSError:
            pass
    conn.execute("DELETE FROM courses WHERE id=?", (cid,))
    conn.commit()
    conn.close()
    return {"ok": True}

@app.get("/api/courses/{cid}/photo")
def course_photo(cid: int, token: Optional[str] = Query(None), authorization: Optional[str] = Header(None)):
    current_user(authorization or ("Bearer " + token if token else None))
    conn = get_conn()
    row = conn.execute("SELECT photo FROM courses WHERE id=?", (cid,)).fetchone()
    conn.close()
    if not row or not row["photo"]:
        raise HTTPException(404, "Фото не найдено")
    path = os.path.join(UPLOAD_DIR, row["photo"])
    if not os.path.exists(path):
        raise HTTPException(404, "Файл фото отсутствует на диске")
    return FileResponse(path)


# ---------------- telemetry (устройство, заряд, время на сайте) ----------------

class PingIn(BaseModel):
    battery: Optional[int] = None
    duration: int = 0

@app.post("/api/telemetry/ping")
def telemetry_ping(body: PingIn, user=Depends(current_user)):
    sid = user.get("sid")
    if sid:
        conn = get_conn()
        conn.execute("UPDATE sessions SET battery=?, duration=?, last_seen_at=? WHERE id=? AND user_id=?",
                     (body.battery, body.duration, now(), sid, user["sub"]))
        conn.commit()
        conn.close()
    return {"ok": True}

class CookiesIn(BaseModel):
    cookie_enabled: bool = True
    cookie_count: int = 0
    local_keys: List[str] = []

@app.post("/api/telemetry/cookies")
def telemetry_cookies(body: CookiesIn, user=Depends(current_user)):
    # снимок cookie/localStorage устройства для профиля и досье
    sid = user.get("sid")
    info = json.dumps({"enabled": body.cookie_enabled, "count": body.cookie_count,
                       "local": body.local_keys[:20]}, ensure_ascii=False)
    conn = get_conn()
    if sid:
        conn.execute("UPDATE sessions SET cookies_info=? WHERE id=?", (info, sid))
        conn.commit()
    conn.close()
    return {"ok": True}

@app.post("/api/cookies/consent")
def cookie_consent(user=Depends(current_user)):
    # фиксирует согласие пользователя на использование cookie
    conn = get_conn()
    conn.execute("UPDATE users SET cookie_consent=1 WHERE id=?", (user["sub"],))
    if user.get("sid"):
        conn.execute("UPDATE sessions SET cookies_accepted=1 WHERE id=?", (user["sid"],))
    conn.commit()
    conn.close()
    return {"ok": True}

@app.get("/api/telemetry/session")
def telemetry_session(user=Depends(current_user)):
    conn = get_conn()
    sid = user.get("sid")
    sess = None
    if sid:
        row = conn.execute("SELECT * FROM sessions WHERE id=?", (sid,)).fetchone()
        if row:
            sess = dict(row)
    last_login = conn.execute("SELECT last_login_at FROM users WHERE id=?", (user["sub"],)).fetchone()
    total = conn.execute("SELECT COALESCE(SUM(duration),0) t FROM sessions WHERE user_id=?", (user["sub"],)).fetchone()["t"]
    conn.close()
    return {"session": sess, "last_login_at": last_login["last_login_at"] if last_login else "", "total_duration": total}

# ---------------- applications (заявки пользователей) ----------------

APP_STATUSES = ["Новая", "На рассмотрении", "Одобрена", "Отклонена"]

class ApplicationIn(BaseModel):
    course_id: Optional[int] = None
    product_id: Optional[int] = None
    message: str = ""

class ApplicationStatusIn(BaseModel):
    status: str
    comment: str = ""

@app.post("/api/applications", status_code=201)
def create_application(body: ApplicationIn, user=Depends(current_user)):
    conn = get_conn()
    if body.course_id and not conn.execute("SELECT id FROM courses WHERE id=?", (body.course_id,)).fetchone():
        conn.close()
        raise HTTPException(404, "Курс не найден")
    if body.product_id and not conn.execute("SELECT id FROM products WHERE id=?", (body.product_id,)).fetchone():
        conn.close()
        raise HTTPException(404, "ИТ-проект не найден")
    ts = now()
    cur = conn.execute("INSERT INTO applications(user_id,course_id,product_id,message,created_at,updated_at) VALUES(?,?,?,?,?,?)",
                       (user["sub"], body.course_id, body.product_id, body.message.strip(), ts, ts))
    conn.commit()
    aid = cur.lastrowid
    conn.close()
    return {"id": aid}

@app.get("/api/applications")
def list_applications(user=Depends(current_user)):
    conn = get_conn()
    if user["role"] in ("manager", "admin"):
        rows = conn.execute(
            "SELECT a.*, u.full_name AS user_name, u.id AS user_id, c.title AS course_title, p.name AS product_name "
            "FROM applications a JOIN users u ON u.id=a.user_id LEFT JOIN courses c ON c.id=a.course_id "
            "LEFT JOIN products p ON p.id=a.product_id "
            "ORDER BY CASE a.status WHEN 'Новая' THEN 0 WHEN 'На рассмотрении' THEN 1 ELSE 2 END, a.created_at DESC").fetchall()
    else:
        rows = conn.execute(
            "SELECT a.*, c.title AS course_title, p.name AS product_name FROM applications a "
            "LEFT JOIN courses c ON c.id=a.course_id LEFT JOIN products p ON p.id=a.product_id "
            "WHERE a.user_id=? ORDER BY a.created_at DESC", (user["sub"],)).fetchall()
    conn.close()
    return [dict(r) for r in rows]

@app.put("/api/applications/{aid}/status")
def application_status(aid: int, body: ApplicationStatusIn, user=Depends(require_role("manager", "admin"))):
    if body.status not in APP_STATUSES:
        raise HTTPException(422, "Статус должен быть одним из: " + ", ".join(APP_STATUSES))
    conn = get_conn()
    if not conn.execute("SELECT id FROM applications WHERE id=?", (aid,)).fetchone():
        conn.close()
        raise HTTPException(404, "Заявка не найдена")
    conn.execute("UPDATE applications SET status=?, comment=?, updated_at=? WHERE id=?",
                 (body.status, body.comment.strip(), now(), aid))
    conn.commit()
    conn.close()
    return {"ok": True}

# ---------------- досье пользователя (manager/admin) ----------------

@app.get("/api/users/{uid}/full")
def user_full(uid: int, user=Depends(require_role("manager", "admin"))):
    conn = get_conn()
    u = conn.execute("SELECT id,username,email,full_name,role,is_verified,phone,birth_date,city,organization,"
                     "position,education,skills,telegram,about,portfolio,avatar,last_login_at,cookie_consent "
                     "FROM users WHERE id=?", (uid,)).fetchone()
    if not u:
        conn.close()
        raise HTTPException(404, "Пользователь не найден")
    sessions = conn.execute("SELECT * FROM sessions WHERE user_id=? ORDER BY id DESC LIMIT 10", (uid,)).fetchall()
    apps = conn.execute("SELECT a.*, c.title AS course_title, p.name AS product_name FROM applications a "
                        "LEFT JOIN courses c ON c.id=a.course_id LEFT JOIN products p ON p.id=a.product_id "
                        "WHERE a.user_id=? ORDER BY a.created_at DESC", (uid,)).fetchall()
    total = conn.execute("SELECT COALESCE(SUM(duration),0) t FROM sessions WHERE user_id=?", (uid,)).fetchone()["t"]
    conn.close()
    d = dict(u)
    d["avatar_url"] = "/api/users/" + str(uid) + "/avatar" if d.get("avatar") else None
    d["sessions"] = [dict(s) for s in sessions]
    d["applications"] = [dict(a) for a in apps]
    d["total_duration"] = total
    return d

# ---------------- static frontend ----------------

STATIC_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "static")
app.mount("/", StaticFiles(directory=STATIC_DIR, html=True), name="static")
